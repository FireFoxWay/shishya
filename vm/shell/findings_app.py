#!/usr/bin/env python3
"""
Findings — what is wrong with your data, without being asked.

Everything else here is reactive: you ask, it answers. This is the opposite.
It reads whatever spreadsheets you have given it and goes looking for the
things nobody thought to ask about — the same way the duplicate supplier
turned up, which is still the most useful thing this project has produced
and which nobody requested.

── Why each check is a generated program ───────────────────────────────────

The model is unreliable at arithmetic across rows and exact at writing code
that does the arithmetic. So a check is a short, bounded brief — "find near
duplicate values in text columns" — turned into a program, run, and reported
with the program attached. Every finding is therefore checkable by the person
being told about it, which is the whole difference between a tool you can
adopt and one you have to trust.

Bounded rather than open-ended for the same reason the chat decides once per
turn: a model this size asked to "analyse this data" wanders. Asked six
specific questions, it answers six specific questions.
"""

import subprocess
import sys
import tempfile
import threading
import time
from pathlib import Path

import gi
gi.require_version("Gtk", "3.0")
gi.require_version("Pango", "1.0")
from gi.repository import GLib, Gtk, Pango   # noqa: E402

import common          # noqa: E402
from chat_app import describe_material, repair, usable, CODE_SYSTEM   # noqa: E402

RUN_TIMEOUT = 90

# Each check is a title, a brief, and why a human should care. The briefs are
# written to be answerable by a program rather than by judgement.
PROTOCOL = (
    "\n\nFirst line of output must be exactly 'FINDINGS: n' where n is how "
    "many problems were found — 0 if none. Then the details, one per line, "
    "each readable by someone who has never seen the file: name the column "
    "and the values in plain text, not as Python tuples or dict reprs.\n\n"
    "Never print more than 10 detail lines. A finding is only useful if it is "
    "short and someone can act on it — a check that reported 8,254 outliers "
    "told the reader nothing at all. If more than 10 things qualify, report "
    "the 10 that matter most and say how many there were in total."
)

CHECKS = [
    ("The same name written two ways",
     "IMPORTANT: where columns repeat as a family — names differing only by a "
     "trailing number, such as 'Purchased From 1' through 'Purchased From 10' "
     "— stack them into ONE series and treat it as a single column. The same "
     "supplier is written one way in one column and another way in the next, "
     "so grouping each column separately finds nothing. "
     "Then, for each text column or stacked family, group values by their "
     "lowercased, whitespace-"
     "stripped form and report only groups where two or more different "
     "spellings occur AND each spelling appears in at least 5 rows. Rank by "
     "the total rows affected, largest first, and report at most 8. Print the "
     "column, the spellings, and the row count of each. Ignore columns where "
     "values are shorter than 4 characters — those are codes and sizes, not "
     "names." + PROTOCOL,
     "The same supplier or manufacturer counted twice splits every total that "
     "involves it."),

    ("Columns nobody fills in",
     "For every column, count rows that are empty, None, zero, 'N/A' or blank. "
     "Report columns above 80%, ranked worst first, at most 10, with the "
     "percentage." + PROTOCOL,
     "A column nobody fills is a broken process or a field that should go."),

    ("Figures that look like typing mistakes",
     "For each numeric column, take only the non-zero values. Skip the column "
     "entirely if it has fewer than 30 non-zero values or if its "
     "interquartile range is zero. Report values more than 20 times the "
     "median, at most 3 per column and at most 10 overall, ranked by how far "
     "out they are, each with a label from the first text column." + PROTOCOL,
     "One misplaced decimal in a price distorts every figure computed from it."),

    ("Identifiers that are not unique",
     "Consider only columns where more than 80 percent of values are distinct "
     "— a genuine identifier, not a size or a category. In those, report "
     "values appearing in more than one row, at most 10, ranked by count." + PROTOCOL,
     "A repeated identifier means two records that should be one."),

    ("Stock that sold and then stopped",
     "If column headers include dates, treat them as a time series per row. "
     "Find rows with total activity above 10 in the earlier half and exactly "
     "zero in the most recent quarter. Rank by how much they used to sell, "
     "largest first, and report at most 10 with a label and the earlier "
     "total. If there are no date columns, say so." + PROTOCOL,
     "Something that sold well and then stopped is either a supply problem or "
     "money asleep on a shelf."),

    ("Value sitting in things that never move",
     "If there is a quantity-on-hand column and a price or cost column, "
     "consider only rows whose dated activity columns are all zero. Print the "
     "total quantity times price for those rows, then the ten largest "
     "individually with a label. If the columns do not exist, say so." + PROTOCOL,
     "The one number that says how much working capital is asleep."),
]


# ─── the checks that do not need a model ────────────────────────────────────
# Where the algorithm is exact and known, writing it is better than asking for
# it. The near-duplicate check is the one that found a supplier counted twice,
# and the model failed it twice running with increasingly clever pandas. A
# check that matters must not depend on the model having a good day.

def check_duplicate_names():
    """Values that are the same name written two ways.

    Repeated column families are stacked before grouping. The same supplier
    appears as "KK ENTERPRISES" in one column and "Kk Enterprises" in the
    next, so grouping each column on its own finds nothing at all — which is
    exactly what happened.
    """
    import re
    from collections import Counter, defaultdict
    from openpyxl import load_workbook

    lines = []
    total = 0
    for path in sorted(common.TRAINING.iterdir()):
        if path.suffix.lower() not in (".xlsx", ".xlsm"):
            continue
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception:
            continue
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                header = [str(c) if c is not None else "" for c in next(rows)]
            except StopIteration:
                continue

            # Columns sharing a stem, minus a trailing number, are one series.
            family = defaultdict(list)
            for i, name in enumerate(header):
                family[re.sub(r"\s*\d+$", "", name).strip() or f"col{i}"].append(i)

            counts = defaultdict(Counter)
            for row in rows:
                for stem, idxs in family.items():
                    for i in idxs:
                        if i >= len(row):
                            continue
                        value = row[i]
                        if not isinstance(value, str):
                            continue
                        value = value.strip()
                        # Short values are sizes and codes, not names.
                        if len(value) < 4 or value.upper() == "N/A":
                            continue
                        counts[stem][value] += 1
            wb_name = path.name
            for stem, tally in counts.items():
                grouped = defaultdict(list)
                for value, n in tally.items():
                    grouped[value.lower()].append((value, n))
                for variants in grouped.values():
                    if len(variants) < 2:
                        continue
                    if not all(n >= 5 for _, n in variants):
                        continue
                    total += 1
                    variants.sort(key=lambda v: -v[1])
                    shown = "  vs  ".join(f"{v} ({n} rows)" for v, n in variants)
                    lines.append((sum(n for _, n in variants),
                                  f"{stem} — {shown}"))
        wb.close()

    lines.sort(reverse=True)
    out = [f"FINDINGS: {total}"]
    out += [text for _, text in lines[:10]]
    if total > 10:
        out.append(f"...and {total - 10} more")
    return "\n".join(out)


BUILTIN = {"The same name written two ways": check_duplicate_names}


class FindingsApp(Gtk.Window):
    def __init__(self):
        super().__init__(title="Findings")
        self.set_default_size(880, 720)
        self.connect("destroy", Gtk.main_quit)
        self.running = False

        root = Gtk.Box(orientation=Gtk.Orientation.VERTICAL)
        self.add(root)

        bar = Gtk.Box(spacing=10, margin=14)
        title = Gtk.Label(xalign=0)
        title.set_markup("<b>Findings</b>")
        bar.pack_start(title, False, False, 0)
        self.status = Gtk.Label(xalign=0)
        self.status.set_opacity(.68)
        self.status.set_text("Looks through your spreadsheets for things nobody asked about.")
        bar.pack_start(self.status, True, True, 0)
        self.go = Gtk.Button(label="Look through my data")
        self.go.get_accessible().set_name("Look through my data")
        self.go.connect("clicked", self.on_go)
        bar.pack_end(self.go, False, False, 0)
        root.pack_start(bar, False, False, 0)

        self.progress = Gtk.ProgressBar()
        self.progress.set_margin_start(14)
        self.progress.set_margin_end(14)
        self.progress.set_show_text(True)
        root.pack_start(self.progress, False, False, 0)

        self.body = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=12, margin=14)
        scroll = Gtk.ScrolledWindow()
        scroll.set_policy(Gtk.PolicyType.NEVER, Gtk.PolicyType.AUTOMATIC)
        scroll.add(self.body)
        self.scroll = scroll
        root.pack_start(scroll, True, True, 0)

    # ── running the checks ──────────────────────────────────────────────────

    def on_go(self, _b):
        if self.running:
            return
        self.running = True
        self.go.set_sensitive(False)
        for child in self.body.get_children():
            self.body.remove(child)
        common.set_state("working")
        threading.Thread(target=self.work, daemon=True).start()

    def work(self):
        listing = describe_material()
        found = 0
        for i, (title, brief, why) in enumerate(CHECKS, 1):
            GLib.idle_add(self.progress.set_fraction, (i - 1) / len(CHECKS))
            GLib.idle_add(self.progress.set_text, f"{i} of {len(CHECKS)} — {title}")
            GLib.idle_add(self.status.set_text, f"Checking: {title.lower()}")

            builtin = BUILTIN.get(title)
            try:
                if builtin:
                    code, out, err = "(built in, not generated)", builtin(), ""
                else:
                    code = repair(common.ask(
                        f"{listing}\n\nWrite a Python script that: {brief}",
                        system=CODE_SYSTEM, budget=5120, job="code"))
                    out, err = ("", "no code") if not usable(code) else run(code)
            except Exception as e:
                code, out, err = "", "", f"{type(e).__name__}: {e}"

            verdict = read_verdict(out)
            if verdict == "nothing":
                GLib.idle_add(self.add_clear, title)
            elif out:
                found += 1
                GLib.idle_add(self.add_finding, title, why, out, code)
            else:
                GLib.idle_add(self.add_failed, title, err)

        GLib.idle_add(self.finish, found)

    def finish(self, found):
        self.progress.set_fraction(1.0)
        self.progress.set_text("done")
        self.status.set_text(
            f"{found} thing{'s' if found != 1 else ''} worth a look."
            if found else "Nothing stood out. That is a good result.")
        self.go.set_sensitive(True)
        self.running = False
        common.set_state("idle")
        return False

    # ── rendering ───────────────────────────────────────────────────────────

    def card(self, accent):
        box = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=5)
        line = Gtk.Box()
        stripe = Gtk.DrawingArea()
        stripe.set_size_request(3, -1)
        stripe.connect("draw", lambda _a, cr: (cr.set_source_rgb(*accent), cr.paint()))
        line.pack_start(stripe, False, False, 0)
        inner = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=5,
                        margin_start=12, margin_top=2, margin_bottom=2)
        line.pack_start(inner, True, True, 0)
        box.pack_start(line, False, False, 0)
        self.body.pack_start(box, False, False, 0)
        return inner

    def add_finding(self, title, why, out, code):
        inner = self.card((0.83, 0.33, 0.13))
        head = Gtk.Label(xalign=0)
        head.set_markup(f"<b>{esc(title)}</b>")
        inner.pack_start(head, False, False, 0)

        reason = Gtk.Label(label=why, xalign=0)
        reason.set_line_wrap(True); reason.set_max_width_chars(88); reason.set_opacity(.68)
        inner.pack_start(reason, False, False, 0)

        result = Gtk.Label(label=out.strip()[:2500], xalign=0, selectable=True)
        result.set_line_wrap(True)
        result.set_line_wrap_mode(Pango.WrapMode.WORD_CHAR)
        result.set_max_width_chars(94)
        result.modify_font(Pango.FontDescription("Monospace 9"))
        inner.pack_start(result, False, False, 0)

        # The program is folded away but present. A finding you cannot check is
        # just an assertion, and this is the difference.
        show = Gtk.Expander(label="Show the program that found this")
        show.get_accessible().set_name("Show the program")
        prog = Gtk.Label(label=code.strip(), xalign=0, selectable=True)
        prog.set_line_wrap(True)
        prog.set_line_wrap_mode(Pango.WrapMode.WORD_CHAR)
        prog.set_max_width_chars(96)
        prog.modify_font(Pango.FontDescription("Monospace 9"))
        prog.set_opacity(.75)
        show.add(prog)
        inner.pack_start(show, False, False, 0)

        self.body.show_all()
        return False

    def add_clear(self, title):
        inner = self.card((0.17, 0.54, 0.24))
        lbl = Gtk.Label(xalign=0)
        lbl.set_markup(f"{esc(title)} — <span alpha='60%'>nothing found</span>")
        inner.pack_start(lbl, False, False, 0)
        self.body.show_all()
        return False

    def add_failed(self, title, err):
        inner = self.card((0.45, 0.45, 0.48))
        lbl = Gtk.Label(xalign=0)
        last = (err or "").strip().splitlines()
        lbl.set_markup(f"{esc(title)} — <span alpha='55%'>could not check: "
                       f"{esc(last[-1][:90] if last else 'no output')}</span>")
        lbl.set_line_wrap(True); lbl.set_max_width_chars(92)
        inner.pack_start(lbl, False, False, 0)
        self.body.show_all()
        return False


# ── helpers ─────────────────────────────────────────────────────────────────

def run(code):
    with tempfile.TemporaryDirectory() as tmp:
        script = Path(tmp) / "check.py"
        script.write_text(code, encoding="utf-8")
        try:
            p = subprocess.run([sys.executable, str(script)],
                               capture_output=True, text=True,
                               timeout=RUN_TIMEOUT, cwd=str(Path.home()))
            return (p.stdout or "").strip(), (p.stderr or "").strip()
        except subprocess.TimeoutExpired:
            return "", f"still running after {RUN_TIMEOUT}s"


def read_verdict(out):
    """How many problems a check found, read from its own first line.

    Inferring this from prose failed immediately: a check that printed "none
    found" once per column produced a wall of nothing that was reported as a
    finding. The script is now asked to say the number itself.
    """
    text = (out or "").strip()
    if not text:
        return "failed"
    first = text.splitlines()[0].strip().lower()
    if first.startswith("findings:"):
        try:
            return "found" if int(first.split(":", 1)[1].strip().split()[0]) > 0 else "nothing"
        except (ValueError, IndexError):
            pass
    # Older wording, and a fallback for a script that ignored the instruction.
    if first in ("none found", "not applicable", "none"):
        return "nothing"
    return "found"


def esc(text):
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


if __name__ == "__main__":
    FindingsApp().show_all()
    Gtk.main()
