#!/usr/bin/env python3
"""
Make a Document — say what you need, get a real file.

Produces genuine .xlsx, .docx and .pptx into Work/ — files that open in Excel,
Word and PowerPoint, not HTML wearing the extension.

── Why the model never writes the file ─────────────────────────────────────

The model returns *content* as JSON — rows, headings, slides — and Python
writes the file. Asking a model to emit a document format directly produces
something that looks right and is subtly corrupt, and a file that will not
open in front of a judge is worse than no feature at all. Structure is the
part a model is good at; serialisation is the part a library is good at.
"""

import json
import re
import threading
import time
from pathlib import Path

import gi
gi.require_version("Gtk", "3.0")
from gi.repository import GLib, Gtk   # noqa: E402

import common   # noqa: E402

# What we ask for back, per kind. Constraining generation to a schema is what
# keeps this from returning prose where a table was wanted.
SCHEMAS = {
    "Spreadsheet": {
        "type": "object",
        "properties": {
            "filename": {"type": "string"},
            "sheets": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "name": {"type": "string"},
                        "columns": {"type": "array", "items": {"type": "string"}},
                        "rows": {"type": "array",
                                 "items": {"type": "array", "items": {"type": "string"}}},
                    },
                    "required": ["name", "columns", "rows"],
                },
            },
        },
        "required": ["filename", "sheets"],
    },
    "Document": {
        "type": "object",
        "properties": {
            "filename": {"type": "string"},
            "title": {"type": "string"},
            "blocks": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "kind": {"type": "string", "enum": ["heading", "paragraph", "bullet"]},
                        "text": {"type": "string"},
                    },
                    "required": ["kind", "text"],
                },
            },
        },
        "required": ["filename", "title", "blocks"],
    },
    "Presentation": {
        "type": "object",
        "properties": {
            "filename": {"type": "string"},
            "title": {"type": "string"},
            "slides": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "heading": {"type": "string"},
                        "bullets": {"type": "array", "items": {"type": "string"}},
                    },
                    "required": ["heading", "bullets"],
                },
            },
        },
        "required": ["filename", "title", "slides"],
    },
}

ASKS = {
    "Spreadsheet": "Design a spreadsheet. Give real column headings and real "
                   "rows of plausible data — never placeholders like 'Row 1'.",
    "Document": "Write the document. Use headings, paragraphs and bullets as "
                "the content actually calls for.",
    "Presentation": "Build the deck. One idea per slide, three to five short "
                    "bullets each — phrases, not sentences.",
}


class MakeApp(Gtk.Window):
    def __init__(self):
        super().__init__(title="Make a Document")
        self.set_default_size(720, 520)
        self.connect("destroy", Gtk.main_quit)

        root = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=0)
        self.add(root)

        head = Gtk.Label(xalign=0, margin=14)
        head.set_markup("<b>Make a Document</b>")
        root.pack_start(head, False, False, 0)

        kinds = Gtk.Box(spacing=8, margin_start=14, margin_end=14)
        self.kind = "Spreadsheet"
        first = None
        for name, hint in (("Spreadsheet", ".xlsx"), ("Document", ".docx"),
                           ("Presentation", ".pptx")):
            btn = Gtk.RadioButton.new_with_label_from_widget(first, f"{name}  {hint}")
            first = first or btn
            btn.get_accessible().set_name(name)
            btn.connect("toggled", self.on_kind, name)
            kinds.pack_start(btn, False, False, 0)
        root.pack_start(kinds, False, False, 0)

        self.brief = Gtk.TextView(wrap_mode=Gtk.WrapMode.WORD, left_margin=10,
                                  right_margin=10, top_margin=8)
        self.brief.get_accessible().set_name("What do you need")
        self.brief.get_buffer().set_text(
            "A register of housing scheme applicants for Warangal district, "
            "with application number, name, income, land holding and status.")
        brief_box = Gtk.ScrolledWindow(margin=14)
        brief_box.set_size_request(-1, 110)
        brief_box.add(self.brief)
        root.pack_start(brief_box, False, False, 0)

        # Anything in Training Data can be used as source material, which is
        # what makes this more than a document generator — it writes *about*
        # the documents already in the sandbox.
        self.use_training = Gtk.CheckButton(label="Use what's in Training Data as source",
                                            margin_start=14)
        self.use_training.get_accessible().set_name("Use Training Data")
        root.pack_start(self.use_training, False, False, 0)

        feet = Gtk.Box(spacing=8, margin=14)
        self.status = Gtk.Label(xalign=0)
        self.status.set_opacity(.7)
        feet.pack_start(self.status, True, True, 0)
        self.make = Gtk.Button(label="Make it")
        self.make.get_accessible().set_name("Make it")
        self.make.connect("clicked", self.on_make)
        feet.pack_end(self.make, False, False, 0)
        root.pack_start(feet, False, False, 0)

        self.log = Gtk.TextView(editable=False, wrap_mode=Gtk.WrapMode.WORD,
                                left_margin=10, right_margin=10)
        self.log.get_accessible().set_name("Result")
        log_box = Gtk.ScrolledWindow(margin=14)
        log_box.add(self.log)
        root.pack_start(log_box, True, True, 0)

    def on_kind(self, btn, name):
        if btn.get_active():
            self.kind = name

    def on_make(self, _b):
        buf = self.brief.get_buffer()
        brief = buf.get_text(buf.get_start_iter(), buf.get_end_iter(), False).strip()
        if not brief:
            return self.say("Say what you need first.")

        self.make.set_sensitive(False)
        self.say("Working it out…")
        common.set_state("working")
        kind = self.kind
        threading.Thread(target=self.work, args=(kind, brief), daemon=True).start()

    def work(self, kind, brief):
        source = ""
        if self.use_training.get_active():
            source = read_training()

        prompt = (f"{ASKS[kind]}\n\nWHAT IS NEEDED:\n{brief}\n"
                  + (f"\nSOURCE MATERIAL — use these facts, do not invent "
                     f"competing ones:\n{source}\n" if source else "")
                  + "\nGive it a short, sensible filename with no extension.")
        try:
            raw = common.ask(prompt, schema=SCHEMAS[kind])
            spec = json.loads(raw)
            path = WRITERS[kind](spec)
            GLib.idle_add(self.done, f"Made {path.name}\n\nIt is in Work/. "
                                     f"Open it from the Work folder on the desktop.")
        except (json.JSONDecodeError, KeyError) as e:
            GLib.idle_add(self.done, f"The model returned something unusable: {e}")
        except Exception as e:
            GLib.idle_add(self.done, f"Could not write the file: {e}")

    def done(self, message):
        self.say(message)
        self.make.set_sensitive(True)
        common.set_state("idle")
        return False

    def say(self, text):
        self.status.set_text(text.split("\n")[0][:70])
        self.log.get_buffer().set_text(text)


def read_training(budget=6000):
    out = []
    if not common.TRAINING.exists():
        return ""
    for p in sorted(common.TRAINING.iterdir()):
        if p.is_file():
            text, _ = common.extract(p)
            out.append(f"[{p.name}]\n{text}")
    return "\n\n".join(out)[:budget]


# ─── writing the actual files ───────────────────────────────────────────────

def safe_name(spec, extension):
    stem = re.sub(r"[^A-Za-z0-9 _-]", "", str(spec.get("filename", "document"))).strip()
    stem = (stem or "document")[:60]
    common.WORK.mkdir(parents=True, exist_ok=True)
    path = common.WORK / f"{stem}{extension}"
    if path.exists():
        path = common.WORK / f"{stem}-{time.strftime('%H%M%S')}{extension}"
    return path


def write_xlsx(spec):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in spec["sheets"] or [{"name": "Sheet1", "columns": [], "rows": []}]:
        ws = wb.create_sheet(title=str(sheet.get("name", "Sheet"))[:31])
        columns = sheet.get("columns", [])
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
        for row in sheet.get("rows", []):
            ws.append(row)
        # Width from content: a spreadsheet whose columns all read ###### is
        # the first thing anyone notices.
        for i, name in enumerate(columns, 1):
            longest = max([len(str(name))] +
                          [len(str(r[i - 1])) for r in sheet.get("rows", [])
                           if len(r) >= i], default=10)
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(48, longest + 3)
        ws.freeze_panes = "A2"
    path = safe_name(spec, ".xlsx")
    wb.save(path)
    return path


def write_docx(spec):
    from docx import Document
    doc = Document()
    doc.add_heading(str(spec.get("title", "Document")), level=0)
    for block in spec.get("blocks", []):
        kind, text = block.get("kind"), str(block.get("text", ""))
        if kind == "heading":
            doc.add_heading(text, level=1)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    path = safe_name(spec, ".docx")
    doc.save(path)
    return path


def write_pptx(spec):
    from pptx import Presentation
    from pptx.util import Pt
    deck = Presentation()

    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = str(spec.get("title", "Presentation"))
    if len(cover.placeholders) > 1:
        cover.placeholders[1].text = time.strftime("%d %B %Y")

    for slide_spec in spec.get("slides", []):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = str(slide_spec.get("heading", ""))
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(slide_spec.get("bullets", [])):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = str(bullet)
            para.font.size = Pt(20)
    path = safe_name(spec, ".pptx")
    deck.save(path)
    return path


WRITERS = {"Spreadsheet": write_xlsx, "Document": write_docx, "Presentation": write_pptx}


if __name__ == "__main__":
    MakeApp().show_all()
    Gtk.main()
