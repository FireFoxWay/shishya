#!/usr/bin/env python3
"""
Code — write a script, or describe one, and run it.

The argument this app makes better than any other: code written by a model
is exactly the code you should not run on your own machine. Here it runs
inside the sandbox — no route to your files, no route to the network — so
you can run it, read the output, and decide. That is the whole thesis with
the stakes turned up.

Highlighting is done with text tags rather than GtkSourceView, which is not
installed. Fewer moving parts, and it only has to be good enough to read.
"""

import re
import subprocess
import sys
import tempfile
import threading
import time
from pathlib import Path

import gi
gi.require_version("Gtk", "3.0")
gi.require_version("Pango", "1.0")
from gi.repository import GLib, Gtk, Pango   # noqa: E402

import common   # noqa: E402

# Generous enough for real work, short enough that a runaway loop does not
# hold the demo hostage. The sandbox would survive either way; your patience
# would not.
RUN_TIMEOUT = 60

STARTER = '''# Anything you write here runs inside the sandbox.
# Your files and the network are out of reach, so it is safe to just run it.

from pathlib import Path

work = Path.home() / "shared" / "work"
for f in sorted(work.iterdir()):
    print(f"{f.name:<44} {f.stat().st_size:>8,} bytes")
'''

KEYWORDS = set("""and as assert async await break class continue def del elif else except
finally for from global if import in is lambda nonlocal not or pass raise return try
while with yield True False None self print len range str int float dict list set""".split())


class CodeApp(Gtk.Window):
    def __init__(self):
        super().__init__(title="Code")
        self.set_default_size(940, 720)
        self.connect("destroy", Gtk.main_quit)
        self.running = None

        root = Gtk.Box(orientation=Gtk.Orientation.VERTICAL)
        self.add(root)

        # ── describe it ──
        top = Gtk.Box(spacing=8, margin=12)
        self.brief = Gtk.Entry(
            placeholder_text="Describe what you want the code to do, then press Write it")
        self.brief.get_accessible().set_name("What should the code do")
        self.brief.connect("activate", self.on_write)
        top.pack_start(self.brief, True, True, 0)

        self.writebtn = Gtk.Button(label="Write it")
        self.writebtn.get_accessible().set_name("Write it")
        self.writebtn.connect("clicked", self.on_write)
        top.pack_end(self.writebtn, False, False, 0)
        root.pack_start(top, False, False, 0)

        # ── editor ──
        self.buffer = Gtk.TextBuffer()
        self.buffer.set_text(STARTER)
        self.editor = Gtk.TextView(buffer=self.buffer, monospace=True,
                                   left_margin=12, right_margin=12, top_margin=10)
        self.editor.get_accessible().set_name("Code")
        self.editor.modify_font(Pango.FontDescription("Monospace 11"))
        self._make_tags()
        self.buffer.connect("changed", self.on_changed)

        editor_scroll = Gtk.ScrolledWindow(margin_start=12, margin_end=12)
        editor_scroll.set_size_request(-1, 360)
        editor_scroll.add(self.editor)
        root.pack_start(editor_scroll, True, True, 0)

        # ── act ──
        bar = Gtk.Box(spacing=8, margin=12)
        self.status = Gtk.Label(xalign=0)
        self.status.set_opacity(.7)
        self.status.set_text("Runs inside the sandbox — nothing here can reach your machine.")
        bar.pack_start(self.status, True, True, 0)
        for label, handler in (("Save to Work", self.on_save), ("Run", self.on_run)):
            b = Gtk.Button(label=label)
            b.get_accessible().set_name(label)
            b.connect("clicked", handler)
            bar.pack_end(b, False, False, 0)
            if label == "Run":
                self.runbtn = b
        root.pack_start(bar, False, False, 0)

        # ── output ──
        self.out = Gtk.TextView(editable=False, monospace=True, left_margin=12,
                                right_margin=12, top_margin=8,
                                wrap_mode=Gtk.WrapMode.WORD_CHAR)
        self.out.get_accessible().set_name("Output")
        out_scroll = Gtk.ScrolledWindow(margin=12)
        out_scroll.set_size_request(-1, 190)
        out_scroll.add(self.out)
        root.pack_start(out_scroll, False, False, 0)

        self.highlight()

    # ── highlighting ────────────────────────────────────────────────────────

    def _make_tags(self):
        t = self.buffer.create_tag
        t("kw", foreground="#0550AE", weight=Pango.Weight.BOLD)
        t("str", foreground="#0A7B47")
        t("com", foreground="#7A7F87", style=Pango.Style.ITALIC)
        t("num", foreground="#A6501E")

    def on_changed(self, _buf):
        # Re-tag on a short delay rather than per keystroke: highlighting a
        # long file on every character makes typing feel sticky.
        if getattr(self, "_pending", None):
            GLib.source_remove(self._pending)
        self._pending = GLib.timeout_add(180, self._rehighlight)

    def _rehighlight(self):
        self._pending = None
        self.highlight()
        return False

    def highlight(self):
        buf = self.buffer
        start, end = buf.get_bounds()
        for name in ("kw", "str", "com", "num"):
            buf.remove_tag_by_name(name, start, end)
        text = buf.get_text(start, end, False)

        def tag(name, a, b):
            buf.apply_tag_by_name(name, buf.get_iter_at_offset(a), buf.get_iter_at_offset(b))

        # Strings and comments first, then skip anything already inside them —
        # otherwise "def" inside a comment lights up as a keyword.
        claimed = []
        for m in re.finditer(r'"""[\s\S]*?"""|\'\'\'[\s\S]*?\'\'\'|"[^"\n]*"|\'[^\'\n]*\'', text):
            tag("str", m.start(), m.end()); claimed.append(m.span())
        for m in re.finditer(r"#[^\n]*", text):
            if not any(a <= m.start() < b for a, b in claimed):
                tag("com", m.start(), m.end()); claimed.append(m.span())

        def free(pos):
            return not any(a <= pos < b for a, b in claimed)

        for m in re.finditer(r"\b\w+\b", text):
            if m.group() in KEYWORDS and free(m.start()):
                tag("kw", m.start(), m.end())
        for m in re.finditer(r"\b\d+\.?\d*\b", text):
            if free(m.start()):
                tag("num", m.start(), m.end())

    # ── the model writes it ─────────────────────────────────────────────────

    SYSTEM = (
        "You write short, complete Python 3 scripts. Output the code and "
        "nothing else — no explanation, no markdown fences. The script runs "
        "in a sandbox on Linux with pandas, openpyxl, python-docx and "
        "python-pptx available. Build paths with pathlib — Path.home() / "
        "'shared' / 'training' / '<name>' — never a string starting with ~, "
        "because that is not expanded and the file will not be found. Print "
        "results rather than returning them, and keep it under 40 lines."
    )

    def on_write(self, _w):
        brief = self.brief.get_text().strip()
        if not brief:
            return self.say("Describe what you want first.")
        self.writebtn.set_sensitive(False)
        self.say("Writing it…")
        common.set_state("working")

        def work():
            try:
                code = strip_fences(common.ask_job("code", 
                    f"{describe_files()}\n\nWrite a Python script that: {brief}",
                    system=self.SYSTEM, budget=6144))
            except RuntimeError as e:
                code = ""
                GLib.idle_add(self.failed, str(e))
                return
            GLib.idle_add(self.wrote, code)

        threading.Thread(target=work, daemon=True).start()

    def failed(self, why):
        self.writebtn.set_sensitive(True)
        common.set_state("idle")
        self.say("Couldn't write that one.")
        self.show(why)
        return False

    def wrote(self, code):
        # Anything that is not code stays out of the editor, or the next press
        # of Run executes an error message and reports a SyntaxError.
        if not code.strip() or code.lstrip().startswith("I ran out of room"):
            return self.failed("The model ran out of room. Try a narrower "
                               "description — one file, one question.")
        self.buffer.set_text(code)
        self.highlight()
        self.writebtn.set_sensitive(True)
        self.say("Read it before you run it.")
        common.set_state("idle")
        return False

    # ── running ─────────────────────────────────────────────────────────────

    def on_run(self, _b):
        if self.running:
            return
        start, end = self.buffer.get_bounds()
        code = self.buffer.get_text(start, end, False)
        if not code.strip():
            return self.say("Nothing to run.")

        self.runbtn.set_sensitive(False)
        self.show("Running…\n")
        self.say("Running…")
        # Running a script competes for the same cores as the model, so the
        # desktop should stand down for this too — not only for generation.
        common.set_state("working")
        self.running = True
        threading.Thread(target=self._run, args=(code,), daemon=True).start()

    def _run(self, code):
        began = time.time()
        with tempfile.TemporaryDirectory() as tmp:
            script = Path(tmp) / "script.py"
            script.write_text(code)
            try:
                p = subprocess.run(
                    [sys.executable, str(script)],
                    capture_output=True, text=True, timeout=RUN_TIMEOUT,
                    cwd=str(common.WORK if common.WORK.exists() else Path.home()),
                )
                out = p.stdout
                if p.stderr:
                    # The traceback is the useful part when something breaks,
                    # so it is shown in full rather than summarised away.
                    out += ("\n" if out else "") + p.stderr
                if not out.strip():
                    out = "(finished, printed nothing)"
                note = f"exit {p.returncode} · {time.time() - began:.1f}s"
            except subprocess.TimeoutExpired:
                out = f"(stopped after {RUN_TIMEOUT}s — it was still going)"
                note = "timed out"
            except Exception as e:
                out, note = f"could not run it: {e}", "failed"
        GLib.idle_add(self.ran, out, note)

    def ran(self, out, note):
        self.show(out)
        self.say(note)
        common.set_state("idle")
        self.runbtn.set_sensitive(True)
        self.running = None
        return False

    # ── saving ──────────────────────────────────────────────────────────────

    def on_save(self, _b):
        start, end = self.buffer.get_bounds()
        common.WORK.mkdir(parents=True, exist_ok=True)
        path = common.WORK / f"script-{time.strftime('%H%M%S')}.py"
        path.write_text(self.buffer.get_text(start, end, False))
        self.say(f"Saved {path.name} to Work")

    def say(self, text):
        self.status.set_text(text)

    def show(self, text):
        self.out.get_buffer().set_text(text)


def describe_files():
    """Tell the model what is actually on disk before it writes anything.

    Left to guess, it invents a plausible filename and plausible column
    positions — asked about the retail inventory it wrote
    "retail_inventory.xlsx" and indexed columns 1 and 3, when the real file
    is "01_20Retail_20Inventory.xlsx" and the columns sit elsewhere. Code
    that reads the wrong file fails in a way that looks like the tool is
    broken, so the filenames and the headers go in the prompt.
    """
    # Training only, and capped. Listing both folders with every column of
    # every spreadsheet ate enough of the context that the model ran out of
    # room mid-thought and produced no code at all.
    lines = []
    for folder, label in ((common.TRAINING, "shared/training"),):
        if not folder.exists():
            continue
        for p in sorted(folder.iterdir())[:8]:
            if not p.is_file() or p.name.startswith("."):
                continue
            entry = f"  ~/{label}/{p.name}"
            head = _headers(p)
            if head:
                entry += f"\n      columns: {head}"
            lines.append(entry)
    if not lines:
        return "There are no files in the sandbox yet."
    return ("These files exist. Use these exact paths and column names — do "
            "not invent others:\n" + "\n".join(lines))


def _headers(path):
    """First row of a spreadsheet, so the model indexes real columns."""
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        return None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        ws = wb.worksheets[0]
        for row in ws.iter_rows(max_row=1, values_only=True):
            wb.close()
            return ", ".join(str(c) for c in row if c is not None)
        wb.close()
    except Exception:
        return None
    return None


def strip_fences(text):
    """Models wrap code in ``` even when told not to."""
    fenced = re.search(r"```(?:python)?\s*([\s\S]*?)```", text)
    return (fenced.group(1) if fenced else text).strip()


if __name__ == "__main__":
    CodeApp().show_all()
    Gtk.main()
