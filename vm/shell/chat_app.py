#!/usr/bin/env python3
"""
Shishya — the dashboard you talk to.

Everything else here is a tool you operate: you choose Training Data, you
choose Make a Document. This is the opposite. You say what you want in
your own words and it decides what to do with your material — answer from
it, build a spreadsheet, draft a document, put together a deck, or write
and run a program to work something out.

That decision is the whole point of this app. The doing was already built
and tested; what was missing was something that chooses, which is the
difference between a toolbox and a colleague.

── Why it chooses once per turn ────────────────────────────────────────────

A model this size, left to run a multi-step loop, wanders. It calls a tool,
misreads the result, calls another, and three steps later is confidently
doing the wrong job. So each turn is a single constrained choice — the same
schema-shaped generation that already produces valid spreadsheets — followed
by a known code path. Less clever, far more likely to work in front of
somebody who matters.
"""

import json
import re
import subprocess
import sys
import tempfile
import threading
import time
from pathlib import Path

import gi
gi.require_version("Gtk", "3.0")
gi.require_version("Gdk", "3.0")
gi.require_version("Pango", "1.0")
from gi.repository import Gdk, GLib, Gtk, Pango   # noqa: E402

import common          # noqa: E402
import make_app        # noqa: E402
from training_app import Library, for_retrieval, SYSTEM as ANSWER_SYSTEM   # noqa: E402

RUN_TIMEOUT = 60

DECIDE = {
    "type": "object",
    "properties": {
        "action": {"type": "string",
                   "enum": ["answer", "spreadsheet", "document", "deck", "code", "clarify"]},
        "because": {"type": "string"},
        "brief": {"type": "string"},
    },
    "required": ["action", "because", "brief"],
}

DECIDE_SYSTEM = """You decide what to do with a request about the user's own documents.

Choose exactly one:
  answer      — any question about rules, policy, eligibility, or what a
                particular record says. This is the default for questions.
                Use it even when the question describes a specific case the
                documents do not name — working out what the rules say about
                a case IS answering.
  spreadsheet — they asked for a table or data file (.xlsx)
  document    — they asked for prose: a note, letter, report, brief (.docx)
  deck        — they asked for slides (.pptx)
  code        — the answer is a number that has to be counted, totalled,
                filtered or ranked across many rows. You are unreliable doing
                that by eye, so write a program instead. Looking up one record
                is not this — that is answer.
  clarify     — you genuinely cannot tell what they are asking for. Almost
                never the right choice. If you could make a reasonable
                attempt, make it. Never use clarify merely because the
                documents might not contain the answer; try, and say so if
                they do not.

"because" is one short sentence to the user saying what you are about to do.
"brief" restates the job precisely for whoever carries it out.
Answer as JSON only."""


CODE_SYSTEM = (
    "You write short, complete Python 3 scripts. Output code only — no "
    "explanation, no markdown fences. Python 3 on Linux with pandas, "
    "openpyxl, python-docx and python-pptx available. Build paths with "
    "pathlib — Path.home() / 'shared' / 'training' / '<name>' — never a "
    "string starting with ~. Print the result. Under 40 lines."
)


class ChatApp(Gtk.Window):
    def __init__(self):
        super().__init__(title="Shishya")
        self.set_default_size(880, 760)
        self.connect("destroy", Gtk.main_quit)
        self.library = None
        self.busy = False
        self.history = []

        root = Gtk.Box(orientation=Gtk.Orientation.VERTICAL)
        self.add(root)

        bar = Gtk.Box(spacing=10, margin=12)
        title = Gtk.Label(xalign=0)
        title.set_markup("<b>Shishya</b>")
        bar.pack_start(title, False, False, 0)
        self.subtitle = Gtk.Label(xalign=0)
        self.subtitle.set_opacity(.65)
        bar.pack_start(self.subtitle, True, True, 0)
        add = Gtk.Button(label="Add sources")
        add.get_accessible().set_name("Add sources")
        add.connect("clicked", self.on_add_sources)
        bar.pack_end(add, False, False, 0)

        reload_btn = Gtk.Button(label="Reread")
        reload_btn.get_accessible().set_name("Reread")
        reload_btn.connect("clicked", lambda _b: self.read_material())
        bar.pack_end(reload_btn, False, False, 0)
        root.pack_start(bar, False, False, 0)

        # ── conversation ──
        self.thread = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=10,
                              margin=14)
        scroll = Gtk.ScrolledWindow()
        scroll.set_policy(Gtk.PolicyType.NEVER, Gtk.PolicyType.AUTOMATIC)
        scroll.add(self.thread)
        self.scroll = scroll
        root.pack_start(scroll, True, True, 0)

        # ── say something ──
        row = Gtk.Box(spacing=8, margin=12)
        self.entry = Gtk.Entry(
            placeholder_text="Tell me what you need — in your own words")
        self.entry.get_accessible().set_name("Message")
        self.entry.connect("activate", self.on_send)
        row.pack_start(self.entry, True, True, 0)
        self.send = Gtk.Button(label="Send")
        self.send.get_accessible().set_name("Send")
        self.send.connect("clicked", self.on_send)
        row.pack_end(self.send, False, False, 0)
        root.pack_start(row, False, False, 0)

        # Drop on the whole window rather than a strip: people aim at the
        # conversation, and a target two pixels wide reads as broken.
        self.drag_dest_set(Gtk.DestDefaults.ALL, [], Gdk.DragAction.COPY)
        self.drag_dest_add_uri_targets()
        self.connect("drag-data-received", self.on_drop)

        # Files also arrive from the control panel, so the folder is watched
        # rather than waiting to be told. Whoever put the file there should
        # not have to remember to announce it.
        self._seen = self.material_fingerprint()
        GLib.timeout_add_seconds(3, self.watch_material)

        self.read_material()
        self.say("shishya",
                 "I've read everything in your Training Data folder. Tell me what "
                 "you need and I'll work out how to make it — an answer, a "
                 "spreadsheet, a document, a deck, or a program if the answer has "
                 "to be counted out of the data.")

    # ── material ────────────────────────────────────────────────────────────

    def read_material(self):
        self.library = Library()
        names = []
        common.TRAINING.mkdir(parents=True, exist_ok=True)
        for p in sorted(common.TRAINING.iterdir()):
            if p.is_file() and not p.name.startswith("."):
                self.library.add(p)
                names.append(p.name)
        self.files = names
        self.subtitle.set_text(
            f"{len(names)} documents · {len(self.library.chunks)} pieces"
            if names else "no material yet — put files in Training Data")

    def material_fingerprint(self):
        try:
            return {p.name: p.stat().st_mtime for p in common.TRAINING.iterdir()
                    if p.is_file() and not p.name.startswith(".")}
        except OSError:
            return {}

    def watch_material(self):
        now = self.material_fingerprint()
        if now != self._seen:
            added = sorted(set(now) - set(self._seen))
            self._seen = now
            self.read_material()
            if added:
                self.say("shishya", self.describe_added(added))
        return True

    def describe_added(self, added):
        names = ", ".join(added)
        return (f"I've read {names}. That's {len(self.files)} documents now — "
                f"ask me anything about them.")

    def on_add_sources(self, _btn):
        chooser = Gtk.FileChooserDialog(
            title="Add sources", transient_for=self, action=Gtk.FileChooserAction.OPEN)
        chooser.add_buttons("Cancel", Gtk.ResponseType.CANCEL, "Add", Gtk.ResponseType.OK)
        chooser.set_select_multiple(True)
        if chooser.run() == Gtk.ResponseType.OK:
            self.take(chooser.get_filenames())
        chooser.destroy()

    def on_drop(self, _w, _ctx, _x, _y, data, _info, _time):
        paths = []
        for uri in data.get_uris():
            p = uri[7:] if uri.startswith("file://") else uri
            paths.append(GLib.uri_unescape_string(p, None) or p)
        self.take(paths)

    def take(self, paths):
        """Copy files in and read them.

        Copied rather than referenced: material has to sit inside the sandbox
        for the clone to reach it at all, and a source that vanishes because
        somebody moved the original is worse than one that takes a little
        disk.
        """
        import shutil
        common.TRAINING.mkdir(parents=True, exist_ok=True)
        added = []
        for raw in paths:
            src = Path(raw)
            if not src.is_file():
                continue
            target = common.TRAINING / src.name
            if src.resolve() != target.resolve():
                shutil.copy2(src, target)
            added.append(src.name)
        if not added:
            return
        self._seen = self.material_fingerprint()
        self.read_material()
        self.say("shishya", self.describe_added(added))

    # ── conversation view ───────────────────────────────────────────────────

    def say(self, who, text, kind=""):
        box = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=3)
        label = Gtk.Label(xalign=0)
        label.set_markup(f"<small><b>{'You' if who == 'you' else 'Shishya'}</b></small>")
        label.set_opacity(.55)
        box.pack_start(label, False, False, 0)

        body = Gtk.Label(xalign=0, label=text)
        # Named so the panel can read replies back without mistaking the
        # user's own message — every bubble is a label otherwise.
        body.get_accessible().set_name("reply" if who != "you" else "yours")
        body.set_line_wrap(True)
        body.set_line_wrap_mode(Pango.WrapMode.WORD_CHAR)
        body.set_selectable(True)
        body.set_max_width_chars(88)
        if kind == "doing":
            body.set_opacity(.7)
        box.pack_start(body, False, False, 0)

        if who == "you":
            box.set_halign(Gtk.Align.END)
        self.thread.pack_start(box, False, False, 0)
        self.thread.show_all()
        GLib.idle_add(self._to_bottom)
        return body

    def _to_bottom(self):
        adj = self.scroll.get_vadjustment()
        adj.set_value(adj.get_upper())
        return False

    # ── the turn ────────────────────────────────────────────────────────────

    def on_send(self, _w):
        if self.busy:
            return
        text = self.entry.get_text().strip()
        if not text:
            return
        self.entry.set_text("")
        self.say("you", text)
        self.busy = True
        self.send.set_sensitive(False)
        common.set_state("working")
        self.working = self.say("shishya", "Thinking about how to do that…", "doing")
        threading.Thread(target=self.work, args=(text,), daemon=True).start()

    def work(self, request):
        try:
            plan = self.decide(request)
        except Exception as e:
            return GLib.idle_add(self.finish, f"I couldn't work out how to do that: {e}")

        GLib.idle_add(self.update, plan.get("because", "Working on it…"))
        action = plan.get("action", "answer")
        brief = plan.get("brief") or request

        try:
            if action == "clarify":
                reply = brief
            elif action == "code":
                reply = self.by_code(brief)
            elif action in ("spreadsheet", "document", "deck"):
                reply = self.by_document(action, brief)
            else:
                reply = self.by_answer(brief)
        except Exception as e:
            reply = f"That didn't work: {type(e).__name__}: {e}"

        GLib.idle_add(self.finish, reply)

    # ── the obvious cases, decided without asking ───────────────────────────
    # The model is good at reading a request and bad at holding a six-way
    # choice steady. Asked twice whether "can a student with 68% attendance
    # sit the exam" was a question, it said clarify once and spreadsheet
    # once — a question it answers well, routed to a refusal and then to a
    # file. So the shapes that are unambiguous in English are settled in
    # English, and the model is asked only about the rest.

    # The second line is money vocabulary. "Find the product contributing good
    # sales with a good margin" contains no counting word at all, so it fell
    # through to the model's six-way choice — the one thing it holds badly —
    # for a question that is arithmetic across 17,011 rows and nothing else.
    # These are the words a pharmacy actually uses to ask for a number.
    COUNTING = ("how many", "how much", "count ", "total ", "sum ", "average",
                "how often", "per category", "each category", "breakdown",
                "most ", "least ", "highest", "lowest", "rank", "top ",
                "margin", "profit", "contribut", "revenue", "turnover",
                "sales value", "best ", "worst ", "growth", "moving",
                "best sell", "top sell", "slowest", "fastest")
    ASKING = ("can ", "could ", "is ", "are ", "does ", "do ", "did ", "should ",
              "what ", "which ", "who ", "when ", "why ", "where ", "how ")
    MAKING = ("make", "build", "create", "draft", "write", "produce", "prepare",
              "give me", "put together")

    def shortcut(self, request):
        """A settled action, or None to let the model choose."""
        q = request.lower().strip()
        wants_file = any(v in q for v in self.MAKING)

        # Counting beats everything: it is the one thing to never do by eye.
        if any(w in q for w in self.COUNTING) and not wants_file:
            return {"action": "code",
                    "because": "That needs counting across the data, so I'll write "
                               "a program to do it exactly.",
                    "brief": request}

        if not wants_file and (q.endswith("?") or q.startswith(self.ASKING)):
            return {"action": "answer",
                    "because": "Let me check your documents.",
                    "brief": request}
        return None

    def decide(self, request):
        settled = self.shortcut(request)
        if settled:
            return settled
        recent = "\n".join(f"- {h}" for h in self.history[-3:])
        prompt = (f"The user's documents are: {', '.join(self.files) or 'none yet'}\n"
                  + (f"Recently:\n{recent}\n" if recent else "")
                  + f"\nRequest: {request}")
        # Every ask here names its job, so the model chosen in the Models
        # app is the one that answers. Without a job name common.ask falls
        # back to the small default, and this window — the front door, the
        # one a judge actually uses — quietly ran all four of its jobs on a
        # 5GB reasoning model while the Models screen showed the 14GB one
        # the owner had picked. Every other app already did this.
        return json.loads(common.ask_job("answers", prompt, system=DECIDE_SYSTEM,
                                         schema=DECIDE, budget=1024))

    # ── the four ways of doing it ───────────────────────────────────────────

    def by_answer(self, brief):
        # Search in English, answer in the language it was asked in. The
        # index is terms; a question in another script matches nothing it
        # should and something it should not.
        query, translated = for_retrieval(brief)
        found = self.library.search(query, k=4)
        if not found:
            return "I can't find anything about that in your material."
        passages = "\n\n".join(f"[from {s}]\n{t}" for s, t in found)
        # The ORIGINAL question, always — the translation was for finding the
        # documents, and the reply must follow the words the person used.
        answer = common.ask_job("answers", f"PASSAGES:\n{passages}\n\nQUESTION: {brief}",
                                system=ANSWER_SYSTEM)
        sources = sorted({s.split(" — ")[0] for s, _ in found})
        self.history.append(f"answered: {brief}")
        return f"{answer.strip()}\n\nRead from: {', '.join(sources)}"

    def by_document(self, action, brief):
        kind = {"spreadsheet": "Spreadsheet", "document": "Document",
                "deck": "Presentation"}[action]
        found = self.library.search(brief, k=10)
        passages = "\n\n".join(f"[from {s}]\n{t}" for s, t in found)
        prompt = (f"{make_app.ASKS[kind]}\n\nWHAT IS NEEDED:\n{brief}\n\n"
                  f"SOURCE MATERIAL — build from these facts and invent no "
                  f"others:\n{passages}\n\nGive it a short filename, no extension.")
        spec = json.loads(common.ask_job("documents", prompt,
                                         schema=make_app.SCHEMAS[kind], budget=4096))
        path = make_app.WRITERS[kind](spec)
        self.history.append(f"made {path.name}")
        return (f"Made {path.name} — it's in your Work folder.\n"
                f"Built from: {', '.join(sorted({s.split(' — ')[0] for s, _ in found}))}")

    def by_code(self, brief):
        """Counting, totalling and ranking go through a program.

        The model is unreliable doing arithmetic across rows by eye — it
        misses rows and occasionally invents a figure. Code that reads the
        file is exact, and the code is on screen to be checked.

        It gets one attempt to fix its own mistakes. The slips are the kind a
        person makes and spots instantly: `idxmin` written without its
        brackets, so pandas was handed the method instead of the number and
        died with an unhashable DataFrame. The reasoning was right and one
        character was wrong. Showing it the traceback is what a person would
        do, and it is far more likely to work than any amount of instruction
        about brackets in advance.
        """
        listing = describe_material()
        code = repair(common.ask_job(
            "code", f"{listing}\n\nWrite a Python script that: {brief}",
            system=CODE_SYSTEM, budget=6144))
        if not usable(code):
            return "I couldn't write a program for that. Try narrowing it down."

        out, err = run_script(code)
        attempts = 1

        if err and not out:
            GLib.idle_add(self.update,
                          "The program hit an error — reading it and trying again.")
            fixed = repair(common.ask_job(
                "code",
                f"{listing}\n\nThis script was meant to: {brief}\n\n"
                f"It failed. Here it is:\n\n{code}\n\n"
                f"The error:\n{err[-1500:]}\n\n"
                f"Return the corrected script, complete, and nothing else.",
                system=CODE_SYSTEM, budget=6144))
            if usable(fixed):
                out2, err2 = run_script(fixed)
                attempts = 2
                if out2:
                    code, out, err = fixed, out2, err2

        self.history.append(f"computed: {brief}")
        note = "" if attempts == 1 else "\n(the first attempt failed; this is the corrected one)"

        if err and not out:
            return (f"I wrote a program twice and it still failed:\n\n{code}\n\n{err}")
        return (f"I worked it out with a program:\n\n{out}\n\n"
                f"— here is what I ran, so you can check it —{note}\n\n{code}")

    # ── finishing ───────────────────────────────────────────────────────────

    def update(self, text):
        self.working.set_text(text)
        return False

    def finish(self, reply):
        self.say("shishya", reply)
        self.busy = False
        self.send.set_sensitive(True)
        common.set_state("idle")
        self.entry.grab_focus()
        return False


def _header_row(ws, look=14):
    """Which row holds the column names. Not always the first.

    A sheet built for a person to read opens with a title, a note about what
    it is, and a block of assumptions; the columns start ten rows down. Taking
    row 1 regardless named this workbook's margin columns "PRODUCT-WISE GROSS
    MARGIN & MOVEMENT ANALYSIS" and eighteen blanks, so every question about
    margin was answered against a sheet whose real columns were never seen.

    The header is the widest row near the top — a title row has one cell, a
    header has nineteen. Ties go to the earliest, which keeps a genuine
    one-column sheet working.
    """
    best, widest = 1, 0
    for n, row in enumerate(ws.iter_rows(min_row=1, max_row=look, values_only=True), 1):
        filled = sum(1 for c in row if c is not None and str(c).strip())
        if filled > widest:
            best, widest = n, filled
    return best


def _describe_sheet(ws, max_rows, max_cols):
    """One sheet's columns, and the values a short column actually contains.

    Names alone are not enough. Told the column was called "Stock Status",
    the model filtered for "Low" when every row says "Low Stock", and
    reported that nothing was low on stock across a sheet with 149 such
    rows — wrong, confident, and impossible to spot without opening the file.
    Columns with few distinct values get those values listed; a column of
    1,000 different prices obviously does not.
    """
    at = _header_row(ws)
    rows = ws.iter_rows(min_row=at, values_only=True)
    try:
        raw_header = [c for c in next(rows) if c is not None]
    except StopIteration:
        return None
    if not raw_header:
        return None
    header = [str(c) for c in raw_header][:max_cols]

    # Which headers are real dates rather than text. Printed as strings they
    # look like text — "2025-04-01 00:00:00" — and the model kept writing
    # .startswith('2025') against them, which fails with "datetime.datetime
    # object has no attribute startswith". Four of eight test questions died
    # on exactly that. Saying so plainly fixes all of them.
    import datetime as _dt
    dated = [str(c) for c in raw_header if isinstance(c, (_dt.datetime, _dt.date))]

    seen = [set() for _ in header]
    counted = 0
    for row in rows:
        if counted >= max_rows:
            break
        counted += 1
        for i, cell in enumerate(row[:len(header)]):
            if cell is not None and len(seen[i]) <= 13:
                seen[i].add(str(cell))

    out = [f"      columns: {', '.join(header)}"]
    if len(raw_header) > max_cols:
        out.append(f"      ({len(raw_header) - max_cols} further columns not listed)")
    if at > 1:
        out.append(f"      NOTE: the column names are on row {at}, not row 1 — "
                   f"read it with pd.read_excel(..., sheet_name={ws.title!r}, "
                   f"header={at - 1}). Rows above that are a title and notes.")
    if dated:
        out.append(
            f"      NOTE: {len(dated)} of those column headers are real datetime "
            f"objects, not text — {dated[0]} through {dated[-1]}. Select them "
            f"with [c for c in df.columns if isinstance(c, datetime)] after "
            f"`from datetime import datetime`. String methods such as "
            f"startswith or str.contains will raise on them. When filtering "
            f"columns BY NAME, guard first — "
            f"[c for c in df.columns if isinstance(c, str) and c.startswith(...)] "
            f"— or the datetime headers will raise as soon as you reach them.")

    # Describe repeated column families rather than listing them one by one.
    # This sheet keeps suppliers wide — "Purchased From 1".."Purchased From 10"
    # each paired with a "Purchase Qty N" and a "Net Purchase Price N". Left to
    # infer that from a flat list of 54 names, the model wrote nothing at all
    # when asked which suppliers we buy most from. Said plainly, it has a
    # chance: the answer needs those pairs stacked before they can be grouped.
    families = {}
    for name in header:
        stem = re.sub(r"\s*\d+$", "", name)
        if stem != name:
            families.setdefault(stem, []).append(name)
    for stem, cols in families.items():
        if len(cols) > 2:
            out.append(f"      {stem} 1..{len(cols)} — one column per repeat")
    if len(families) > 1 and any(len(c) > 2 for c in families.values()):
        pairs = " and ".join(f'"{stem} N"' for stem in families if len(families[stem]) > 2)
        out.append(f"      these repeat together: {pairs} share the same N, so "
                   f"to total by value across them, stack the pairs first "
                   f"(pd.concat of each N, or pd.wide_to_long) and then group")

    for name, values in zip(header, seen):
        if 1 < len(values) <= 12:
            out.append(f"      {name} contains only: {', '.join(sorted(values))}")
    return "\n".join(out)


def sheet_shape(path, max_rows=400, max_sheets=16, max_cols=44):
    """Every sheet in the workbook, not only the first.

    A one-sheet export was the only shape this had ever been given. A real
    workbook is a stack of sheets that refer to each other — a master list,
    and a dozen views built on top of it — and describing only the first left
    the model computing from raw columns while a sheet answering the exact
    question sat three tabs away, unmentioned.
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    blocks = []
    for ws in wb.worksheets[:max_sheets]:
        try:
            block = _describe_sheet(ws, max_rows, max_cols)
        except Exception:
            continue          # one unreadable sheet must not hide the rest
        if block:
            blocks.append(f"    sheet {ws.title!r}:\n{block}")
    extra = len(wb.worksheets) - max_sheets
    wb.close()
    if not blocks:
        return None
    if extra > 0:
        blocks.append(f"    ({extra} further sheets not described)")
    return "\n".join(blocks)


# The listing goes into every code prompt, so it cannot grow without limit —
# one fourteen-sheet workbook took it past the context window on its own. The
# window is now sized to fit (see common.ask), but a prompt that grows with the
# training folder is one that will eventually be mostly file listing.
#
# Trimmed, never dropped. The first version skipped any file that would not fit
# whole, which threw away the largest workbook — the only one that could answer
# the question — while keeping four small files nobody had asked about. A path
# always goes in, because the model cannot read a file it was not told about;
# only the sheet detail is cut, at a line boundary, and it says so.
LISTING_BUDGET = 16_000
NOTE_BUDGET = 3_500


def describe_billing():
    """The shop's own books, described so a program can read them.

    Kept out of the material folder deliberately: training material is
    reference — a policy, a price list, a manual — and these two files are
    live operational data that changed while you were reading this. They are
    named separately in the prompt for the same reason, so the model is never
    guessing whether a row is history or today.
    """
    import csv as _csv
    out = []
    for path, what in ((common.CATALOGUE, "what is stocked right now, with cost, "
                                          "price, stock on hand and reorder point"),
                       (common.SALES, "every line item ever sold, append-only, one "
                                      "row per product per bill")):
        if not path.exists():
            continue
        try:
            with open(path, newline="", encoding="utf-8") as f:
                reader = _csv.reader(f)
                header = next(reader, [])
                rows = sum(1 for _ in reader)
        except OSError:
            continue
        out.append(f"  {path}\n"
                   f"      {what}\n"
                   f"      columns: {', '.join(header)}   ({rows:,} rows)")
    if not out:
        return ""
    return ("\n\nTHE SHOP'S LIVE BOOKS — written by the Billing app, current as "
            "of now. `at` is 'YYYY-MM-DD HH:MM:SS' text; parse with "
            "pd.to_datetime. `profit` is already (price − cost) × qty at the "
            "cost on the day of sale — sum it, never recompute margin from "
            "today's catalogue cost:\n" + "\n".join(out))


def describe_material():
    lines, spent = [], 0
    for p in sorted(common.TRAINING.iterdir()):
        if not p.is_file() or p.name.startswith("."):
            continue
        entry = f"  {common.TRAINING}/{p.name}"
        spent += len(entry)
        shape = None
        if p.suffix.lower() in (".xlsx", ".xlsm"):
            try:
                shape = sheet_shape(p)
            except Exception:
                shape = None
        if shape:
            room = LISTING_BUDGET - spent
            if len(shape) > room:
                kept = shape[:max(room, 0)].rsplit("\n", 1)[0]
                shape = (kept + "\n    (further sheets not described)") if kept else None
            if shape:
                entry += "\n" + shape
                spent += len(shape)
        lines.append(entry)
    if not lines:
        return "There are no files yet."
    listing = ("These files exist — use these exact paths and column names:\n"
               + "\n".join(lines))

    # Written guidance reaches the code path too, not only the answering one.
    #
    # A note in the material saying "this sheet costs every sale at the
    # cheapest of ten purchase prices, so its margins are three times too
    # good" is exactly what someone writing a program needs, and it was only
    # ever retrieved for prose questions. The program went on computing from
    # a column the note exists to warn about. A file named *-notes.md is the
    # convention: material about the material.
    # Every notes file that fits, not the first two. A count cap meant the
    # third file written was silently never read — and nothing says so.
    notes, spent = [], 0
    for p in sorted(common.TRAINING.glob("*notes.md")):
        try:
            text = p.read_text(encoding="utf-8")[:NOTE_BUDGET]
        except OSError:
            continue
        if spent + len(text) > 3 * NOTE_BUDGET:
            break
        notes.append(f"--- {p.name} ---\n{text}")
        spent += len(text)
    if notes:
        listing += ("\n\nRead these notes before writing the program. They "
                    "describe traps in the data above and override anything a "
                    "column name implies:\n" + "\n".join(notes))
    listing += describe_billing()
    return listing


def repair_regex(code):
    r"""Make string literals containing regex escapes raw.

    Asked which supplier gives the best average price, it reached for
    pd.wide_to_long with suffix='\d+' — no r prefix — which Python reads as
    an invalid escape and pandas then fails on. The intent is unambiguous
    every time: a backslash followed by a regex class letter inside a short
    literal is a pattern, not text.
    """
    import re

    def raw(m):
        quote, body = m.group(1), m.group(2)
        return f"r{quote}{body}{quote}"

    # Short literals only, and only those carrying a regex escape. A long
    # string with a stray backslash is more likely to be prose.
    return re.sub(r"(?<![rbRB])(['\"])((?:[^'\"\\]|\\.){0,40}?\\[dDwWsSbBAZ][^'\"]{0,20}?)\1",
                  raw, code)


def repair(raw):
    """Everything we fix about generated code, in one place."""
    return repair_regex(repair_calls(repair_paths(strip_fences(raw))))


def usable(code):
    """Is this a script, or the model's apology?"""
    return bool(code.strip()) and not code.lstrip().startswith(
        ("I ran out of room", "I thought my way", "I couldn't"))


def run_script(code):
    """Run it in the sandbox and return what it said. Errors come back whole —
    the traceback is the useful part when the model has to read it back."""
    with tempfile.TemporaryDirectory() as tmp:
        script = Path(tmp) / "job.py"
        script.write_text(code, encoding="utf-8")
        try:
            p = subprocess.run([sys.executable, str(script)],
                               capture_output=True, text=True,
                               timeout=RUN_TIMEOUT, cwd=str(Path.home()))
            return (p.stdout or "").strip(), (p.stderr or "").strip()
        except subprocess.TimeoutExpired:
            return "", f"it was still running after {RUN_TIMEOUT}s"


# Methods this model writes without their brackets. Shown the traceback it
# made the identical slip a second time, so this is repaired mechanically
# rather than explained again: `df.loc[df['x'].idxmin]` hands pandas the
# method instead of the number and dies with "unhashable type: DataFrame".
#
# Only methods. Properties — size, shape, values, empty, columns, index,
# dtypes, T — are correct without brackets and adding them would break code
# that was already right.
NEEDS_BRACKETS = (
    "idxmin", "idxmax", "mean", "median", "nunique", "unique", "tolist",
    "to_dict", "dropna", "drop_duplicates", "reset_index", "value_counts",
    "sort_values", "sort_index", "count", "sum", "min", "max", "abs",
    "round", "first_valid_index", "last_valid_index", "keys", "item",
)


def repair_calls(code):
    """Add the brackets the model left off a method call."""
    import re
    pattern = r"\.(" + "|".join(NEEDS_BRACKETS) + r")\b(?!\s*[\(\w=])"
    return re.sub(pattern, r".\1()", code)


def repair_paths(code):
    """Point filenames at files that exist.

    Given "01_20Retail_20Inventory.xlsx" the model wrote
    "01_20Retail_Inventory.xlsx" — one fragment dropped. The script then fails
    on a missing file, which reads as the whole tool being broken rather than
    as a typo. Any quoted name ending in a document extension that does not
    exist is matched against the real ones and replaced when there is a clear
    best match.
    """
    import difflib
    import re

    real = [p.name for p in common.TRAINING.iterdir() if p.is_file()]
    real += [p.name for p in common.WORK.iterdir() if p.is_file()] if common.WORK.exists() else []
    if not real:
        return code

    def swap(m):
        quote, name = m.group(1), m.group(2)
        if name in real:
            return m.group(0)
        close = difflib.get_close_matches(name, real, n=1, cutoff=0.6)
        return f"{quote}{close[0]}{quote}" if close else m.group(0)

    return re.sub(r"(['\"])([^'\"/]+\.(?:xlsx|xlsm|csv|docx|pptx|pdf|txt|md))\1",
                  swap, code)


def strip_fences(text):
    import re
    m = re.search(r"```(?:python)?\s*([\s\S]*?)```", text)
    return (m.group(1) if m else text).strip()


if __name__ == "__main__":
    ChatApp().show_all()
    Gtk.main()
